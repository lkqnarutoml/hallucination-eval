#!/usr/bin/env python3
"""Normalize agent trajectories into conversations.json, then split them into batches.

Subcommands
    inspect   FILE [--sheet NAME] [--out column_map.proposed.json]
              Propose a column map for an xlsx/csv export: headers, masked samples, confidence per field.
    normalize INPUT [INPUT ...] --out DIR [--column-map map.json] [--format auto|events_jsonl|ces_json|xlsx|csv]
              [--ground-truth FILE] [--sheet NAME]
              Write DIR/conversations.json and DIR/normalize_report.md.
    split     CONVERSATIONS_JSON --out DIR [--max-conversations 12] [--max-chars 40000]
              Write DIR/batch_NNN.json files and DIR/index.json.

Exit codes: 0 ok; 2 unreadable input or missing dependency; 3 a tabular export needs a confirmed column
map (run `inspect`, confirm the proposal with the user, then pass --column-map).

The normalized record keeps the ORDER of what happened inside a turn (utterance, tool call, state write)
because grounding checks depend on it, and every event carries a locator back to the source line or cell.
Standard library only; openpyxl is imported lazily for .xlsx files.
"""
from __future__ import annotations

import argparse
import copy
import csv
import datetime as _dt
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from halu_common.jsonl import read_json, read_jsonl, write_json  # noqa: E402
from halu_common.locators import column_letter, csv_locator, events_locator, json_locator, xlsx_locator  # noqa: E402
from halu_common.payload import canonical_tool_name, find_key, parse_payload  # noqa: E402
from halu_common.redact import cap_payload, redact_text, redact_value, truncate_string  # noqa: E402
from halu_common.textnum import mask_digits  # noqa: E402

SKILL_VERSION = "0.4.0"
REQUIRED_FIELDS = ("session_id", "text", "tool_name", "tool_args", "tool_response")
SYNONYMS: Dict[str, List[str]] = {
    "session_id": ["session_id", "sessionid", "conversation_id", "conversationid", "conversation", "call_id",
                   "interaction_id", "interactionid", "session", "trace_id", "dialog_id", "dialogue_id"],
    "timestamp": ["timestamp", "event_time", "eventtime", "time", "created_at", "publish_time", "ts", "datetime",
                  "date_time", "event_timestamp", "logged_at"],
    "turn": ["turn", "turn_index", "turn_id", "turn_number", "step", "turn_no"],
    "row_type": ["row_type", "event_type", "eventtype", "type", "kind", "event", "record_type", "entry_type"],
    "speaker": ["speaker", "role", "author", "participant", "source", "actor", "sender", "party"],
    "text": ["text", "utterance", "message", "content", "transcript", "query", "response_text", "user_text",
             "agent_text", "utterance_text"],
    "tool_name": ["tool_name", "toolname", "function_name", "functionname", "tool", "name", "displayname",
                  "display_name", "function"],
    "tool_args": ["tool_args", "toolargs", "arguments", "args", "input", "tool_input", "parameters", "request",
                  "function_args", "tool_arguments", "tool_parameters"],
    "tool_response": ["tool_response", "toolresponse", "output", "tool_output", "result", "response",
                      "function_response", "tool_result", "return_value"],
    "tool_call_id": ["tool_call_id", "call_id", "toolcallid", "function_call_id", "invocation_id"],
    "state": ["state", "variables", "updated_variables", "updatedvariables", "session_state", "context_data",
              "session_variables", "state_delta"],
    "ces_json": ["diagnostic_info", "diagnosticinfo", "raw_json", "raw", "payload_json", "ces_response",
                 "raw_response", "trace"],
    "agent": ["agent", "agent_name", "agentname", "current_agent"],
    "flow": ["flow", "issue_reason", "workflow", "intent", "caller_intent"],
}
DEFAULT_VALUES = {
    "customer_speakers": ["user", "customer", "caller", "human", "end_user", "client"],
    "agent_speaker_regex": r"(?i)agent|democard|bot|model|assistant|system|virtual",
    "tool_row_types": ["tool_call", "tool", "function_call", "function", "toolcall", "tool_response", "tool_result"],
    "utterance_row_types": ["utterance", "text", "transcript", "message", "user", "agent", "customer", "assistant"],
    "state_row_types": ["state", "updatedvariables", "updated_variables", "variable_update", "variables"],
    "end_row_types": ["end", "transfer", "hangup", "session_end", "disconnect", "sys.remote-call-disconnected"],
    "abandon_markers": ["sys.remote-call-disconnected"],
}
END_TOOL_NAMES = ("democard_end_conversation", "end_session", "end_conversation", "transfer_to_human")
_ROLE_CUSTOMER_RE = re.compile(r"(?i)^(user|customer|caller|human|end_user|client)$")
_UUID_RE = re.compile(r"^[0-9a-f]{8}-?[0-9a-f]{4}-?[0-9a-f]{4}-?[0-9a-f]{4}-?[0-9a-f]{12}$", re.I)
_IDENT_RE = re.compile(r"^[A-Za-z_][\w\-./]*$")


# ----------------------------------------------------------------------------------------------------------
# small utilities
# ----------------------------------------------------------------------------------------------------------
def now_iso() -> str:
    return _dt.datetime.now(_dt.timezone.utc).replace(microsecond=0).isoformat()


def parse_ts(value: Any) -> Optional[_dt.datetime]:
    if value is None or value == "":
        return None
    if isinstance(value, _dt.datetime):
        return value if value.tzinfo else value.replace(tzinfo=_dt.timezone.utc)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            seconds = float(value)
            if seconds > 1e12:
                seconds /= 1000.0
            return _dt.datetime.fromtimestamp(seconds, tz=_dt.timezone.utc)
        except (OverflowError, OSError, ValueError):
            return None
    text = str(value).strip()
    if not text:
        return None
    candidate = text.replace("Z", "+00:00").replace("z", "+00:00")
    if " " in candidate and "T" not in candidate:
        candidate = candidate.replace(" ", "T", 1)
    match = re.match(r"^(.*\.\d{1,6})\d+(.*)$", candidate)  # trim nanoseconds
    if match:
        candidate = match.group(1) + match.group(2)
    try:
        parsed = _dt.datetime.fromisoformat(candidate)
    except ValueError:
        return None
    return parsed if parsed.tzinfo else parsed.replace(tzinfo=_dt.timezone.utc)


def norm_header(name: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(name or "").lower())


def flatten(value: Any, prefix: str = "") -> Dict[str, Any]:
    if isinstance(value, dict):
        out: Dict[str, Any] = {}
        for key, item in value.items():
            out.update(flatten(item, "%s.%s" % (prefix, key) if prefix else str(key)))
        return out
    return {prefix: value}


def diff_state(before: dict, after: dict) -> Dict[str, dict]:
    old, new = flatten(before), flatten(after)
    changes: Dict[str, dict] = {}
    for path in sorted(old.keys() | new.keys()):
        if path in old and path not in new:
            changes[path] = {"old": old[path], "new": None, "deleted": True}
        elif path not in old:
            changes[path] = {"old": None, "new": new[path]}
        elif old[path] != new[path]:
            changes[path] = {"old": old[path], "new": new[path]}
    return changes


def compact_state(state: dict, max_chars: int = 500) -> dict:
    """Scalars verbatim, small containers verbatim, large containers summarized."""
    out = {}
    for key, value in state.items():
        if isinstance(value, (dict, list)):
            encoded = json.dumps(value, ensure_ascii=False, default=str)
            out[key] = value if len(encoded) <= max_chars else {"_summary": "%s with %d items" % (type(value).__name__, len(value))}
        else:
            out[key] = value
    return out


def new_conversation(conversation_id: str, source_path: str, fmt: str, locator: str) -> dict:
    return {
        "conversation_id": str(conversation_id),
        "source": {"path": source_path, "format": fmt, "locator": locator},
        "metadata": {"scenario_id": None, "mode": None, "input_modality": None, "flow": None,
                     "ground_truth": None, "ground_truth_source": None},
        "events": [],
        "warnings": [],
    }


def add_event(conv: dict, kind: str, turn: Optional[int], **fields: Any) -> dict:
    event = {"index": len(conv["events"]), "turn": turn, "kind": kind}
    event.update(fields)
    conv["events"].append(event)
    return event


def warn(conv: dict, code: str, detail: str = "", locator: Optional[str] = None) -> None:
    conv["warnings"].append({"code": code, "detail": detail, "locator": locator})


def outcome_from_reason(reason: Any) -> str:
    """Map a session_end reason (runner stop reason, telephony marker) to completed / transferred / abandoned."""
    text = str(reason or "").lower()
    if not text:
        return "unknown"
    if any(k in text for k in ("transfer", "escalat", "handoff", "human", "representative")):
        return "transferred"
    if any(k in text for k in ("submitted", "completed", "success", "hangup", "provided", "end_signal")):
        return "completed"
    if any(k in text for k in ("turn_limit", "simulator", "timeout", "disconnect", "abandon", "error", "refused")):
        return "abandoned"
    return "unknown"


def finalize(conv: dict) -> dict:
    events = conv["events"]
    turn = 0
    for event in events:
        if event.get("turn") is None:
            if event["kind"] == "customer_utterance":
                turn += 1
            event["turn"] = turn
    tool_calls = [e for e in events if e["kind"] == "tool_call"]
    outcome = "unknown"
    for call in tool_calls:
        if call.get("tool") in END_TOOL_NAMES:
            args = call.get("args") or {}
            escalated, _ = find_key(args, ("session_escalated", "escalated"))
            outcome = "transferred" if str(escalated).lower() == "true" else "completed"
    if outcome == "unknown":
        for event in events:
            if event["kind"] == "session_end":
                outcome = outcome_from_reason(event.get("reason"))
    last_state = {}
    for event in events:
        if event["kind"] == "state_update" and isinstance(event.get("state"), dict):
            last_state = event["state"]
    flow, _ = find_key(last_state, ("issue_reason", "flow", "workflow"))
    if flow and not conv["metadata"].get("flow"):
        conv["metadata"]["flow"] = flow
    conv["summary"] = {
        "turns": max([e["turn"] for e in events] + [0]),
        "events": len(events),
        "tool_calls": len(tool_calls),
        "tools": sorted({c.get("tool") for c in tool_calls if c.get("tool")}),
        "outcome": outcome,
        "customer_utterances": sum(1 for e in events if e["kind"] == "customer_utterance"),
        "agent_utterances": sum(1 for e in events if e["kind"] == "agent_utterance"),
    }
    return conv


# ----------------------------------------------------------------------------------------------------------
# CES diagnosticInfo chunks (the representation that keeps intra-turn order)
# ----------------------------------------------------------------------------------------------------------
def steps_from_messages(response: Any, state: dict, locator_for) -> List[dict]:
    """Turn-ordered steps from response.outputs[*].diagnosticInfo.messages[*].chunks. Mutates ``state``."""
    steps: List[dict] = []
    pending: Dict[Any, dict] = {}
    outputs = (response or {}).get("outputs", []) if isinstance(response, dict) else []
    for oi, out in enumerate(outputs or []):
        diag = (out or {}).get("diagnosticInfo") or {}
        for mi, message in enumerate(diag.get("messages", []) or []):
            role = message.get("role", "") or ""
            when = message.get("eventTime")
            for ci, chunk in enumerate(message.get("chunks", []) or []):
                if not isinstance(chunk, dict):
                    continue
                pointer = "/outputs/%d/diagnosticInfo/messages/%d/chunks/%d" % (oi, mi, ci)
                locator = locator_for(pointer)
                if "text" in chunk or "transcript" in chunk:
                    is_transcript = "text" not in chunk
                    text = chunk.get("transcript") if is_transcript else chunk.get("text")
                    if text is None:
                        continue
                    kind = "customer_utterance" if role == "user" else "agent_utterance"
                    if steps and steps[-1]["kind"] == kind and steps[-1].get("_fragment") and is_transcript:
                        steps[-1]["text"] += str(text)
                        continue
                    step = {"kind": kind, "text": str(text), "time": when, "locator": locator, "_fragment": is_transcript,
                            "text_source": "ces_transcript" if is_transcript else "ces_text"}
                    if kind == "agent_utterance":
                        step["agent"] = role
                    steps.append(step)
                elif "toolCall" in chunk:
                    call = chunk.get("toolCall") or {}
                    step = {"kind": "tool_call", "tool": canonical_tool_name(call.get("tool"), call.get("displayName")),
                            "raw_tool": call.get("displayName") or call.get("tool"), "args": call.get("args", {}),
                            "response": None, "call_id": call.get("id"), "agent": role, "time": when, "locator": locator}
                    steps.append(step)
                    pending[call.get("id")] = step
                elif "toolResponse" in chunk:
                    reply = chunk.get("toolResponse") or {}
                    step = pending.pop(reply.get("id"), None)
                    if step is not None:
                        step["response"] = reply.get("response")
                        step["response_locator"] = locator
                    else:
                        steps.append({"kind": "tool_call", "tool": canonical_tool_name(reply.get("tool"), reply.get("displayName")),
                                      "raw_tool": reply.get("displayName") or reply.get("tool"), "args": None,
                                      "response": reply.get("response"), "call_id": reply.get("id"), "agent": role,
                                      "time": when, "locator": locator, "unmatched_response": True})
                elif "agentTransfer" in chunk:
                    transfer = chunk.get("agentTransfer") or {}
                    steps.append({"kind": "transfer", "target": transfer.get("displayName") or canonical_tool_name(transfer.get("targetAgent")),
                                  "agent": role, "time": when, "locator": locator})
                elif "defaultVariables" in chunk:
                    before = copy.deepcopy(state)
                    for key, value in (chunk.get("defaultVariables") or {}).items():
                        state.setdefault(key, copy.deepcopy(value))
                    changes = diff_state(before, state)
                    if changes:  # the initial state, visible once so that evaluators do not need snapshots
                        steps.append({"kind": "state_update", "source": "defaults", "changes": changes, "state": compact_state(state),
                                      "time": when, "locator": locator})
                elif "updatedVariables" in chunk:
                    before = copy.deepcopy(state)
                    state.update(copy.deepcopy(chunk.get("updatedVariables") or {}))
                    changes = diff_state(before, state)
                    if changes:
                        steps.append({"kind": "state_update", "source": "updatedVariables", "changes": changes, "state": compact_state(state),
                                      "time": when, "locator": locator})
                elif "deletedVariables" in chunk:
                    before = copy.deepcopy(state)
                    for key in chunk.get("deletedVariables") or []:
                        state.pop(key, None)
                    changes = diff_state(before, state)
                    if changes:
                        steps.append({"kind": "state_update", "source": "deletedVariables", "changes": changes, "state": compact_state(state),
                                      "time": when, "locator": locator})
    for step in steps:
        step.pop("_fragment", None)
    return steps


def finish_tool_step(step: dict) -> dict:
    """Parse args/response payloads of a tool step and record their representations."""
    args = parse_payload(step.get("args"))
    response = parse_payload(step.get("response"))
    step["args"] = args.as_dict() if args.representation != "empty" else {}
    step["response"] = response.as_dict() if response.representation != "empty" else {}
    step["representation"] = {"args": args.representation, "response": response.representation}
    if args.meta:
        step.setdefault("meta", {}).update(args.meta)
    return step


# ----------------------------------------------------------------------------------------------------------
# events.jsonl (runner / simulator captures)
# ----------------------------------------------------------------------------------------------------------
def discover_events_files(path: Path, max_depth: int = 3) -> List[Path]:
    if path.is_file():
        return [path]
    found: List[Path] = []
    for candidate in sorted(path.rglob("events.jsonl")):
        depth = len(candidate.relative_to(path).parts) - 1
        if depth <= max_depth:
            found.append(candidate)
    return found


def _origin_lookup(turn_events: List[dict]):
    """Return (by_call_id, by_time) origin classifiers built from the turn's execution spans."""
    spans = []
    for event in turn_events:
        if event.get("event") in ("execution_model", "execution_callback", "execution_tool"):
            span = event.get("span") or {}
            spans.append((parse_ts(span.get("start")), event["event"], span.get("attributes") or {}))
    spans = [s for s in spans if s[0] is not None]
    spans.sort(key=lambda item: item[0])

    def _origin_before(index: int) -> str:
        for start, kind, attrs in reversed(spans[:index]):
            if kind == "execution_model":
                return "model"
            if kind == "execution_callback":
                stage = str(attrs.get("stage", ""))
                return "model" if stage == "AfterModel" else "callback"
        return "unknown"

    by_call_id: Dict[str, str] = {}
    for index, (start, kind, attrs) in enumerate(spans):
        if kind == "execution_tool":
            call_id = attrs.get("tool call id") or attrs.get("tool_call_id")
            if call_id:
                by_call_id[call_id] = _origin_before(index)

    def by_time(when: Optional[str]) -> str:
        moment = parse_ts(when)
        if moment is None or not spans:
            return "unknown"
        index = 0
        while index < len(spans) and spans[index][0] <= moment:
            index += 1
        return _origin_before(index)

    return by_call_id, by_time


def load_events_file(path: Path) -> List[dict]:
    events = read_jsonl(path)
    by_session: Dict[str, List[dict]] = defaultdict(list)
    run_level: List[dict] = []
    for event in events:
        if not isinstance(event, dict) or "event" not in event:
            continue
        sid = event.get("session_id")
        if sid:
            by_session[str(sid)].append(event)
        else:
            run_level.append(event)
    if not by_session:
        # captures without session ids: group by run id (or the folder) so nothing is lost
        for event in run_level:
            by_session[str(event.get("run_id") or path.parent.name)].append(event)
        run_level = []
    conversations = []
    for sid, evs in by_session.items():
        evs.sort(key=lambda e: e.get("sequence", 0))
        first_seq = evs[0].get("sequence", 0)
        conv = new_conversation(sid, str(path), "events_jsonl", events_locator(str(path), first_seq))
        meta = conv["metadata"]
        for event in evs:
            if event.get("event") == "session_start":
                record = event.get("test_record") or {}
                scenario = event.get("scenario") or {}
                meta["scenario_id"] = event.get("scenario_id") or record.get("id") or scenario.get("id")
                meta["mode"] = event.get("mode")
                truth = {}
                if record.get("zip_code_answer"):
                    truth["zip_on_file"] = str(record["zip_code_answer"])
                facts = scenario.get("facts") or {}
                if facts.get("zip"):
                    truth["zip_on_file"] = str(facts["zip"])
                if truth:
                    meta["ground_truth"] = truth
                    meta["ground_truth_source"] = "fixture"
            if event.get("event") == "customer_message" and meta.get("input_modality") is None:
                meta["input_modality"] = event.get("input_modality")
            if meta.get("scenario_id") is None and event.get("scenario_id"):
                meta["scenario_id"] = event.get("scenario_id")
            if meta.get("mode") is None and event.get("mode"):
                meta["mode"] = event.get("mode")
        by_turn: Dict[Any, List[dict]] = defaultdict(list)
        for event in evs:
            by_turn[event.get("turn")].append(event)
        turns = sorted((t for t in by_turn if t is not None), key=lambda t: (isinstance(t, str), t))
        state: dict = {}
        used_chunks = False
        for turn in turns:
            turn_events = by_turn[turn]
            ces = [e for e in turn_events if e.get("event") == "ces_response"]
            by_call_id, by_time = _origin_lookup(turn_events)
            steps: List[dict] = []
            if ces:
                used_chunks = True
                for event in ces:
                    seq = event.get("sequence", 0)
                    steps.extend(steps_from_messages(event.get("response"), state,
                                                     lambda ptr, _seq=seq: events_locator(str(path), _seq, ptr)))
            if not any(s["kind"] == "customer_utterance" for s in steps):
                for event in turn_events:
                    if event.get("event") == "customer_message":
                        text = event.get("text") or event.get("message") or ""
                        steps.insert(0, {"kind": "customer_utterance", "text": str(text), "time": event.get("timestamp"),
                                         "locator": events_locator(str(path), event.get("sequence", 0)),
                                         "text_source": "customer_message"})
                        break
            if not ces:
                spans = [e for e in turn_events if e.get("event") == "execution_tool"]
                spans.sort(key=lambda e: str((e.get("span") or {}).get("start") or ""))
                for event in spans:
                    attrs = (event.get("span") or {}).get("attributes") or {}
                    steps.append({"kind": "tool_call", "tool": canonical_tool_name(attrs.get("name")), "raw_tool": attrs.get("name"),
                                  "args": attrs.get("args"), "response": attrs.get("response"),
                                  "call_id": attrs.get("tool call id") or attrs.get("tool_call_id"),
                                  "agent": attrs.get("agent") or (event.get("span") or {}).get("agent"),
                                  "time": (event.get("span") or {}).get("start"),
                                  "locator": events_locator(str(path), event.get("sequence", 0))})
                for event in turn_events:
                    if event.get("event") == "state_snapshot" and isinstance(event.get("state"), dict):
                        before = copy.deepcopy(state)
                        state = copy.deepcopy(event["state"])
                        changes = diff_state(before, state)
                        if changes:
                            steps.append({"kind": "state_update", "source": "state_snapshot", "changes": changes, "state": compact_state(state),
                                          "time": event.get("timestamp"),
                                          "locator": events_locator(str(path), event.get("sequence", 0))})
                if not spans:
                    warn(conv, "no_execution_spans", "turn %s has neither ces_response chunks nor execution spans" % turn)
            if not any(s["kind"] == "agent_utterance" for s in steps):
                for event in turn_events:
                    if event.get("event") == "agent_message":
                        text = event.get("text") or event.get("message") or ""
                        if text:
                            steps.append({"kind": "agent_utterance", "text": str(text), "time": event.get("timestamp"),
                                          "locator": events_locator(str(path), event.get("sequence", 0)),
                                          "text_source": "agent_message", "agent": event.get("agent")})
                        break
            for step in steps:
                kind = step.pop("kind")
                if kind == "tool_call":
                    finish_tool_step(step)
                    origin = by_call_id.get(step.get("call_id") or "", None)
                    step["origin"] = origin or by_time(step.get("time"))
                elif kind == "agent_utterance":
                    step["origin"] = by_time(step.get("time"))
                add_event(conv, kind, turn, **step)
        if not turns:
            warn(conv, "no_turns", "no events carried a turn number")
        if not used_chunks:
            warn(conv, "fallback_used", "no ces_response chunks; used execution spans and message events")
        end_event = next((e for e in evs if e.get("event") == "session_end"), None)
        if end_event is not None:
            add_event(conv, "session_end", None, reason=end_event.get("end_reason") or end_event.get("stop_reason"),
                      escalated=None, locator=events_locator(str(path), end_event.get("sequence", 0)))
            if not [e for e in conv["events"] if e["kind"] == "tool_call"] and end_event.get("tool_calls"):
                for call in end_event.get("tool_calls") or []:
                    step = {"tool": canonical_tool_name(call.get("name")), "raw_tool": call.get("name"),
                            "args": call.get("args"), "response": call.get("response"), "call_id": call.get("call_id"),
                            "agent": call.get("agent"), "origin": "unknown",
                            "locator": events_locator(str(path), end_event.get("sequence", 0))}
                    finish_tool_step(step)
                    add_event(conv, "tool_call", None, **step)
                warn(conv, "flat_tool_calls_only", "tool calls taken from session_end.tool_calls (no turn order)")
        conversations.append(finalize(conv))
    return conversations


# ----------------------------------------------------------------------------------------------------------
# raw CES response JSON files (one directory per session, e.g. turn-001-response.json)
# ----------------------------------------------------------------------------------------------------------
def load_ces_json(path: Path) -> List[dict]:
    files = [path] if path.is_file() else sorted(p for p in path.rglob("*.json") if "response" in p.name.lower())
    by_dir: Dict[Path, List[Path]] = defaultdict(list)
    for file in files:
        by_dir[file.parent].append(file)
    conversations = []
    for folder, members in by_dir.items():
        conv = new_conversation(folder.name, str(folder), "ces_json", json_locator(str(members[0]), "/"))
        state: dict = {}
        for turn, file in enumerate(sorted(members), start=1):
            try:
                response = read_json(file)
            except ValueError as error:
                warn(conv, "unreadable_json", "%s: %s" % (file, error))
                continue
            steps = steps_from_messages(response, state, lambda ptr, _f=str(file): json_locator(_f, ptr))
            for step in steps:
                kind = step.pop("kind")
                if kind == "tool_call":
                    finish_tool_step(step)
                    step["origin"] = "unknown"
                add_event(conv, kind, turn, **step)
        conversations.append(finalize(conv))
    return conversations


# ----------------------------------------------------------------------------------------------------------
# tabular exports (xlsx / csv) with a column map
# ----------------------------------------------------------------------------------------------------------
def read_table(path: Path, sheet: Optional[str] = None):
    """Return (header, [(row_number, record)], locator_fn, sheet_name)."""
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            print("error: reading .xlsx needs openpyxl (python -m pip install openpyxl)", file=sys.stderr)
            sys.exit(2)
        import openpyxl as _openpyxl

        book = _openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        worksheet = book[sheet] if sheet and sheet in book.sheetnames else book[book.sheetnames[0]]
        header: Optional[List[str]] = None
        records: List[Tuple[int, dict]] = []
        for number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            values = list(row)
            if header is None:
                if sum(1 for v in values if v not in (None, "")) >= 3:
                    header = [str(v).strip() if v not in (None, "") else "col%d" % (i + 1) for i, v in enumerate(values)]
                continue
            if all(v in (None, "") for v in values):
                continue
            record = {header[i]: values[i] for i in range(min(len(header), len(values)))}
            records.append((number, record))
        if header is None:
            print("error: no header row found in %s" % path, file=sys.stderr)
            sys.exit(2)
        sheet_name = worksheet.title
        header_index = {name: i + 1 for i, name in enumerate(header)}

        def locator(row_number: int, field: Optional[str]) -> str:
            column = column_letter(header_index.get(field or "", 1))
            return xlsx_locator(str(path), sheet_name, "%s%d" % (column, row_number))

        return header, records, locator, sheet_name
    if ext in (".csv", ".tsv", ".txt"):
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            sample = handle.read(16384)
            handle.seek(0)
            if ext == ".tsv":
                dialect: Any = csv.excel_tab
            else:
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                except csv.Error:
                    dialect = csv.excel
            reader = csv.DictReader(handle, dialect=dialect)
            header = [h.strip() if h else "col%d" % (i + 1) for i, h in enumerate(reader.fieldnames or [])]
            reader.fieldnames = header
            records = []
            for row in reader:
                if all(v in (None, "") for v in row.values()):
                    continue
                records.append((reader.line_num, dict(row)))
        if not header:
            print("error: no header row found in %s" % path, file=sys.stderr)
            sys.exit(2)
        return header, records, (lambda row_number, field: csv_locator(str(path), row_number)), None
    print("error: unsupported tabular file %s" % path, file=sys.stderr)
    sys.exit(2)


def _name_score(column: str, names: List[str]) -> float:
    normalized = norm_header(column)
    best = 0.0
    for name in names:
        candidate = norm_header(name)
        if normalized == candidate:
            return 1.0
        if candidate and candidate in normalized:
            best = max(best, 0.7 if len(candidate) >= 4 else 0.4)
    return best


def _ratio(values: List[Any], predicate) -> float:
    present = [v for v in values if v not in (None, "")]
    if not present:
        return 0.0
    return sum(1 for v in present if predicate(v)) / float(len(present))


def _value_score(field: str, values: List[Any]) -> float:
    present = [v for v in values if v not in (None, "")]
    if not present:
        return 0.0
    unique = len({str(v) for v in present})
    cardinality = unique / float(len(present))
    text = [str(v) for v in present]
    if field == "session_id":
        idlike = _ratio(present, lambda v: bool(_UUID_RE.match(str(v))) or len(str(v)) >= 12 and " " not in str(v))
        return idlike * (1.0 if 0.005 <= cardinality <= 0.9 else 0.4)
    if field == "timestamp":
        return _ratio(present, lambda v: parse_ts(v) is not None)
    if field == "turn":
        return _ratio(present, lambda v: str(v).strip().isdigit() and int(str(v).strip()) < 1000)
    if field == "row_type":
        keywords = ("tool", "user", "agent", "utter", "state", "text", "call", "customer", "model")
        return (1.0 if unique <= 20 else 0.2) * _ratio(present, lambda v: any(k in str(v).lower() for k in keywords))
    if field == "speaker":
        keywords = ("user", "agent", "customer", "model", "assistant", "bot", "caller", "human")
        return (1.0 if unique <= 30 else 0.2) * _ratio(present, lambda v: any(k in str(v).lower() for k in keywords))
    if field == "text":
        return _ratio(present, lambda v: " " in str(v) and len(str(v)) >= 8 and not str(v).lstrip().startswith(("{", "[")))
    if field == "tool_name":
        return (1.0 if unique <= 60 else 0.3) * _ratio(present, lambda v: bool(_IDENT_RE.match(str(v))) and " " not in str(v))
    if field == "tool_args":
        return _ratio(present, lambda v: isinstance(parse_payload(v).value, dict) and bool(parse_payload(v).value))
    if field == "tool_response":
        def _looks_response(v: Any) -> bool:
            parsed = parse_payload(v).value
            return isinstance(parsed, dict) and any(k in parsed for k in ("status", "result", "response", "output"))
        return _ratio(present, _looks_response)
    if field == "tool_call_id":
        return (1.0 if cardinality > 0.5 else 0.3) * _ratio(present, lambda v: " " not in str(v) and len(str(v)) >= 8)
    if field == "state":
        return _ratio(present, lambda v: isinstance(parse_payload(v).value, dict) and len(parse_payload(v).value) >= 2)
    if field == "ces_json":
        return _ratio(present, lambda v: "diagnosticInfo" in str(v) or '"chunks"' in str(v))
    if field == "agent":
        return _ratio(present, lambda v: str(v).lower().endswith("agent"))
    if field == "flow":
        return (0.6 if unique <= 10 else 0.0)
    return 0.0


def propose_column_map(header: List[str], records: List[Tuple[int, dict]], sample_size: int = 200) -> dict:
    sample = [record for _, record in records[:sample_size]]
    scored = []  # (score, field, column, evidence)
    for field, names in SYNONYMS.items():
        for column in header:
            name_score = _name_score(column, names)
            values = [record.get(column) for record in sample]
            value_score = _value_score(field, values)
            if name_score > 0:
                score = 0.65 * name_score + 0.35 * value_score
            else:
                score = 0.5 * value_score
            if score >= 0.3:
                scored.append((round(score, 3), field, column, "name %.2f, values %.2f" % (name_score, value_score)))
    scored.sort(key=lambda item: -item[0])
    proposal: Dict[str, dict] = {}
    used_columns = set()
    for score, field, column, evidence in scored:
        if field in proposal or column in used_columns:
            continue
        samples = []
        for record in sample:
            value = record.get(column)
            if value not in (None, "") and len(samples) < 3:
                samples.append(mask_digits(str(value))[:80])
        proposal[field] = {"column": column, "confidence": score, "evidence": evidence, "samples": samples}
        used_columns.add(column)
    missing = [f for f in REQUIRED_FIELDS if f not in proposal or proposal[f]["confidence"] < 0.8]
    return {"headers": header, "rows_sampled": len(sample), "proposal": proposal,
            "required_missing_or_low_confidence": missing, "needs_confirmation": True,
            "suggested_column_map": {"version": 1, "sheet": None,
                                     "columns": {f: proposal[f]["column"] if f in proposal else None for f in SYNONYMS},
                                     "values": DEFAULT_VALUES, "ordering": "timestamp_then_row"}}


def _in_list(value: str, patterns: List[str]) -> bool:
    lowered = value.lower()
    return any(p.lower() == lowered or p.lower() in lowered for p in patterns if p)


def _role_from_row(record: dict, cols: dict, vals: dict) -> Tuple[str, Optional[str]]:
    speaker_col = cols.get("speaker")
    speaker = str(record.get(speaker_col) or "").strip() if speaker_col else ""
    customers = vals.get("customer_speakers", DEFAULT_VALUES["customer_speakers"])
    agent_regex = re.compile(vals.get("agent_speaker_regex", DEFAULT_VALUES["agent_speaker_regex"]))
    if speaker:
        if speaker.lower() in [c.lower() for c in customers] or _ROLE_CUSTOMER_RE.match(speaker):
            return "customer_utterance", None
        if agent_regex.search(speaker) or speaker.lower().endswith("agent"):
            return "agent_utterance", speaker
    row_type_col = cols.get("row_type")
    row_type = str(record.get(row_type_col) or "").lower() if row_type_col else ""
    if any(k in row_type for k in ("user", "customer", "caller", "human")):
        return "customer_utterance", None
    return "agent_utterance", speaker or None


def classify_row(record: dict, cols: dict, vals: dict) -> str:
    ces_col = cols.get("ces_json")
    if ces_col and record.get(ces_col) not in (None, "") and "diagnosticInfo" in str(record.get(ces_col)):
        return "ces_json"
    row_type_col = cols.get("row_type")
    tool = record.get(cols["tool_name"]) if cols.get("tool_name") else None
    text = record.get(cols["text"]) if cols.get("text") else None
    if row_type_col and record.get(row_type_col) not in (None, ""):
        row_type = str(record[row_type_col]).strip()
        if _in_list(row_type, vals.get("tool_row_types", DEFAULT_VALUES["tool_row_types"])):
            return "tool_call"
        if _in_list(row_type, vals.get("state_row_types", DEFAULT_VALUES["state_row_types"])):
            return "state_update"
        if _in_list(row_type, vals.get("end_row_types", DEFAULT_VALUES["end_row_types"])):
            return "session_end"
        if _in_list(row_type, vals.get("utterance_row_types", DEFAULT_VALUES["utterance_row_types"])):
            return "utterance"
    if tool not in (None, ""):
        return "tool_call"
    if text not in (None, ""):
        return "utterance"
    if cols.get("state") and record.get(cols["state"]) not in (None, ""):
        return "state_update"
    return "ignore"


def load_tabular(path: Path, column_map: dict, sheet: Optional[str] = None, fmt: str = "xlsx") -> List[dict]:
    cols = {k: v for k, v in (column_map.get("columns") or {}).items() if v}
    vals = column_map.get("values") or DEFAULT_VALUES
    missing = [f for f in REQUIRED_FIELDS if f not in cols]
    if missing:
        print("error: column map lacks required fields: %s" % ", ".join(missing), file=sys.stderr)
        sys.exit(3)
    header, records, locator, _sheet_name = read_table(path, sheet or column_map.get("sheet"))
    unknown = [c for c in cols.values() if c not in header]
    if unknown:
        print("error: column map names columns that are not in the file: %s" % ", ".join(unknown), file=sys.stderr)
        sys.exit(3)
    sessions: Dict[str, List[Tuple[int, dict]]] = defaultdict(list)
    dropped = 0
    for row_number, record in records:
        sid = str(record.get(cols["session_id"]) or "").strip()
        if not sid:
            dropped += 1
            continue
        sessions[sid].append((row_number, record))
    conversations = []
    ts_col = cols.get("timestamp")
    turn_col = cols.get("turn")
    ordering = column_map.get("ordering", "timestamp_then_row")
    for sid, rows in sessions.items():
        conv = new_conversation(sid, str(path), fmt, locator(rows[0][0], cols["session_id"]))
        if dropped:
            warn(conv, "rows_without_session_id", "%d rows in the file had no session id" % dropped)
            dropped = 0
        ordered = rows
        if ts_col and ordering == "timestamp_then_row":
            stamped = [(parse_ts(record.get(ts_col)), row_number, record) for row_number, record in rows]
            rate = sum(1 for s in stamped if s[0] is not None) / float(len(stamped))
            if rate >= 0.95:
                floor = _dt.datetime.min.replace(tzinfo=_dt.timezone.utc)
                ordered = [(rn, rec) for _, rn, rec in sorted(stamped, key=lambda s: (s[0] or floor, s[1]))]
                if [rn for rn, _ in ordered] != [rn for rn, _ in rows]:
                    warn(conv, "ordering_conflict", "timestamp order differs from file order; timestamp order used")
            else:
                warn(conv, "ordering_file_order", "only %.0f%% of timestamps parsed; file order used" % (rate * 100))
        seen_rows = set()
        state: dict = {}
        turn = 0
        open_calls: Dict[str, dict] = {}
        last_call: Optional[dict] = None
        for row_number, record in ordered:
            signature = json.dumps({k: str(v) for k, v in record.items()}, sort_keys=True)
            if signature in seen_rows:
                warn(conv, "duplicate_row", "row %d duplicates an earlier row" % row_number, locator(row_number, None))
                continue
            seen_rows.add(signature)
            if turn_col and str(record.get(turn_col) or "").strip().isdigit():
                turn = int(str(record[turn_col]).strip())
            kind = classify_row(record, cols, vals)
            timestamp = str(record.get(ts_col)) if ts_col and record.get(ts_col) not in (None, "") else None
            if kind == "ignore":
                continue
            if kind == "ces_json":
                parsed = parse_payload(record.get(cols["ces_json"])).value
                for step in steps_from_messages(parsed, state, lambda ptr, _rn=row_number: locator(_rn, cols["ces_json"]) + "#" + ptr):
                    step_kind = step.pop("kind")
                    if step_kind == "tool_call":
                        finish_tool_step(step)
                        step["origin"] = "unknown"
                    add_event(conv, step_kind, turn or None, **step)
                continue
            if kind == "utterance":
                role, agent = _role_from_row(record, cols, vals)
                if role == "customer_utterance" and not turn_col:
                    turn += 1
                event = {"text": str(record.get(cols["text"]) or ""), "time": timestamp, "locator": locator(row_number, cols["text"]),
                         "text_source": "column:%s" % cols["text"]}
                if role == "agent_utterance":
                    event["agent"] = agent or (str(record.get(cols["agent"])) if cols.get("agent") and record.get(cols["agent"]) else None)
                    event["origin"] = "unknown"
                add_event(conv, role, turn or None, **event)
                continue
            if kind == "tool_call":
                args = parse_payload(record.get(cols["tool_args"]))
                response = parse_payload(record.get(cols["tool_response"]))
                name = canonical_tool_name(record.get(cols["tool_name"]))
                call_id = str(record.get(cols["tool_call_id"])) if cols.get("tool_call_id") and record.get(cols["tool_call_id"]) not in (None, "") else None
                if args.representation == "empty" and response.representation != "empty":
                    target = open_calls.get(call_id or "") or (last_call if last_call and last_call.get("tool") == name and not last_call.get("response") else None)
                    if target is not None:
                        target["response"] = response.as_dict()
                        target["representation"]["response"] = response.representation
                        target["response_locator"] = locator(row_number, cols["tool_response"])
                        continue
                event = {"tool": name, "raw_tool": record.get(cols["tool_name"]), "args": args.as_dict() if args.representation != "empty" else {},
                         "response": response.as_dict() if response.representation != "empty" else {},
                         "representation": {"args": args.representation, "response": response.representation},
                         "call_id": call_id, "agent": str(record.get(cols["agent"])) if cols.get("agent") and record.get(cols["agent"]) else None,
                         "origin": "unknown", "time": timestamp, "locator": locator(row_number, cols["tool_args"])}
                if args.meta:
                    event["meta"] = args.meta
                stored = add_event(conv, "tool_call", turn or None, **event)
                last_call = stored
                if call_id:
                    open_calls[call_id] = stored
                if cols.get("state") and record.get(cols["state"]) not in (None, ""):
                    parsed_state = parse_payload(record.get(cols["state"])).value
                    if isinstance(parsed_state, dict) and parsed_state:
                        before = copy.deepcopy(state)
                        state.update(copy.deepcopy(parsed_state))
                        changes = diff_state(before, state)
                        if changes:
                            add_event(conv, "state_update", turn or None, source="column", changes=changes, state=compact_state(state),
                                      time=timestamp, locator=locator(row_number, cols["state"]))
                continue
            if kind == "state_update":
                parsed = parse_payload(record.get(cols["state"])).value if cols.get("state") else {}
                if isinstance(parsed, dict) and parsed:
                    before = copy.deepcopy(state)
                    state.update(copy.deepcopy(parsed))
                    changes = diff_state(before, state)
                    if changes:
                        add_event(conv, "state_update", turn or None, source="column", changes=changes, state=compact_state(state),
                                  time=timestamp, locator=locator(row_number, cols["state"]))
                continue
            if kind == "session_end":
                marker = str(record.get(cols["row_type"]) if cols.get("row_type") else "")
                add_event(conv, "session_end", turn or None, reason=marker, escalated=None, locator=locator(row_number, None))
                if _in_list(marker, vals.get("abandon_markers", DEFAULT_VALUES["abandon_markers"])):
                    conv["metadata"]["abandoned"] = True
        conversations.append(finalize(conv))
    return conversations


# ----------------------------------------------------------------------------------------------------------
# ground truth attachment
# ----------------------------------------------------------------------------------------------------------
def attach_ground_truth(conversations: List[dict], path: Path) -> int:
    rows: List[dict] = []
    if path.suffix.lower() == ".json":
        data = read_json(path)
        rows = data if isinstance(data, list) else [dict(value, _id=key) for key, value in data.items()]
    else:
        header, records, _, _ = read_table(path)
        rows = [record for _, record in records]
    if not rows:
        return 0
    ids = {c["conversation_id"] for c in conversations}
    id_column = None
    for column in rows[0].keys():
        values = {str(r.get(column)) for r in rows}
        if column == "_id" or norm_header(column) in {norm_header(n) for n in SYNONYMS["session_id"]} or (values & ids):
            id_column = column
            break
    if id_column is None:
        return 0
    by_id = {str(r.get(id_column)): r for r in rows}
    attached = 0
    for conv in conversations:
        row = by_id.get(conv["conversation_id"])
        if row is None:
            continue
        truth = {k: v for k, v in row.items() if k != id_column and v not in (None, "")}
        merged = dict(conv["metadata"].get("ground_truth") or {})
        merged.update(truth)
        conv["metadata"]["ground_truth"] = merged
        conv["metadata"]["ground_truth_source"] = str(path)
        attached += 1
    return attached


# ----------------------------------------------------------------------------------------------------------
# report
# ----------------------------------------------------------------------------------------------------------
DIET_DEFAULTS = {"state_snapshots": "none", "redact": True, "redact_opaque": False, "cap": True, "max_string": 600,
                 "max_depth": 2, "list_items": 8, "max_payload": 2500}


def _cap(value: Any, settings: dict, counts: Dict[str, int]) -> Any:
    if not settings.get("cap", True):
        return value
    if isinstance(value, (dict, list)):
        return cap_payload(value, settings["max_string"], settings["max_depth"], settings["list_items"], settings["max_payload"], counts)
    if isinstance(value, str):
        return truncate_string(value, settings["max_string"], counts)
    return value


def _summarize_old(value: Any, settings: dict, counts: Dict[str, int]) -> Any:
    """The previous value of a changed variable: scalars are kept (redacted, truncated), containers become a summary."""
    if not settings.get("cap", True):
        return value
    if isinstance(value, dict):
        counts["old_summarized"] = counts.get("old_summarized", 0) + 1
        return {"_summary": "dict with %d keys" % len(value)}
    if isinstance(value, list):
        counts["old_summarized"] = counts.get("old_summarized", 0) + 1
        return [{"_summary": "list with %d items" % len(value)}]
    return _cap(value, settings, counts)


def apply_diet(conv: dict, settings: dict, counts: Dict[str, int]) -> dict:
    """Shrink and sanitize one normalized conversation in place, after finalize() has read the raw snapshots.

    state_snapshots: "none" drops the `state` copy of every state_update (evaluators apply `changes`), "turn"
    keeps the last one of each turn, "all" keeps them all. Kept snapshots, every changed value, tool args and
    responses are redacted (unless redact is off) and capped; utterances are redacted, never truncated.
    """
    before = len(json.dumps(conv, ensure_ascii=False, default=str))
    events = conv.get("events", [])
    last_per_turn: Dict[Any, int] = {}
    for index, event in enumerate(events):
        if event.get("kind") == "state_update":
            last_per_turn[event.get("turn")] = index
    redact = bool(settings.get("redact"))
    opaque = bool(settings.get("redact_opaque"))
    for index, event in enumerate(events):
        kind = event.get("kind")
        if kind == "state_update":
            policy = settings.get("state_snapshots", "none")
            keep = policy == "all" or (policy == "turn" and last_per_turn.get(event.get("turn")) == index)
            if "state" in event:
                if keep:
                    state = redact_value(event["state"], None, counts, opaque) if redact else event["state"]
                    event["state"] = _cap(state, settings, counts)
                else:
                    del event["state"]
                    counts["snapshots_dropped"] = counts.get("snapshots_dropped", 0) + 1
            for path, change in (event.get("changes") or {}).items():
                key = path.split(".")[-1]
                for field in ("old", "new"):
                    value = change.get(field)
                    if value is None:
                        continue
                    if redact:
                        value = redact_value(value, key, counts, opaque)
                    change[field] = _cap(value, settings, counts) if field == "new" else _summarize_old(value, settings, counts)
        elif kind == "tool_call":
            for field in ("args", "response"):
                value = event.get(field)
                if value is None:
                    continue
                if redact:
                    value = redact_value(value, None, counts, opaque)
                event[field] = _cap(value, settings, counts)
        elif kind in ("customer_utterance", "agent_utterance") and redact and isinstance(event.get("text"), str):
            event["text"] = redact_text(event["text"], counts)
    after = len(json.dumps(conv, ensure_ascii=False, default=str))
    conv.setdefault("summary", {})["chars_before"] = before
    conv["summary"]["chars_after"] = after
    counts["chars_before"] = counts.get("chars_before", 0) + before
    counts["chars_after"] = counts.get("chars_after", 0) + after
    return conv


def diet_settings_from_args(args: argparse.Namespace) -> dict:
    settings = dict(DIET_DEFAULTS)
    if getattr(args, "no_diet", False):
        return {"state_snapshots": "all", "redact": False, "redact_opaque": False, "cap": False, "max_string": 0, "max_depth": 99, "list_items": 0, "max_payload": 0}
    settings["state_snapshots"] = args.state_snapshots
    settings["redact"] = not args.no_redact
    settings["redact_opaque"] = bool(args.redact_opaque)
    settings["max_string"] = args.max_string
    settings["max_depth"] = args.max_depth
    settings["list_items"] = args.list_items
    settings["max_payload"] = args.max_payload
    return settings


def build_report(sources: List[dict], conversations: List[dict], diet: Optional[dict] = None) -> str:
    kinds = Counter()
    representations = Counter()
    warnings = Counter()
    tool_rows = 0
    parsed_rows = 0
    for conv in conversations:
        for event in conv["events"]:
            kinds[event["kind"]] += 1
            if event["kind"] == "tool_call":
                tool_rows += 1
                rep = event.get("representation") or {}
                representations["args=%s" % rep.get("args")] += 1
                representations["response=%s" % rep.get("response")] += 1
                if rep.get("args") not in ("unparseable",) and rep.get("response") not in ("unparseable",):
                    parsed_rows += 1
        for warning in conv["warnings"]:
            warnings[warning["code"]] += 1
    lines = ["# Normalization report", "", "Generated: %s" % now_iso(), "", "## Sources", "",
             "| Path | Format | Conversations |", "|---|---|---|"]
    for source in sources:
        lines.append("| %s | %s | %d |" % (source["path"], source["format"], source["conversations"]))
    lines += ["", "## Totals", "", "- Conversations: %d" % len(conversations),
              "- Events by kind: %s" % ", ".join("%s=%d" % kv for kv in sorted(kinds.items())),
              "- Tool calls: %d, payload parse rate: %s" % (tool_rows, ("%.1f%%" % (100.0 * parsed_rows / tool_rows)) if tool_rows else "n/a"),
              "- Payload representations seen: %s" % (", ".join("%s=%d" % kv for kv in sorted(representations.items())) or "none"),
              "", "## Warnings", ""]
    if warnings:
        for code, count in sorted(warnings.items()):
            lines.append("- %s: %d" % (code, count))
    else:
        lines.append("- none")
    if diet:
        settings = diet.get("settings") or {}
        counts = diet.get("counts") or {}
        lines += ["", "## Diet (what was removed before anything reaches a model)", "",
                  "- Settings: state_snapshots=%s, redact=%s, redact_opaque=%s, cap=%s, max_string=%s, max_depth=%s, list_items=%s, max_payload=%s" % (
                      settings.get("state_snapshots"), settings.get("redact"), settings.get("redact_opaque"), settings.get("cap", True), settings.get("max_string"),
                      settings.get("max_depth"), settings.get("list_items"), settings.get("max_payload")),
                  "- Characters: %s before -> %s after" % (counts.get("chars_before", 0), counts.get("chars_after", 0)),
                  "", "| Rule | Count |", "|---|---|"]
        rule_rows = [(k, v) for k, v in sorted(counts.items()) if k not in ("chars_before", "chars_after")]
        lines += ["| %s | %d |" % (k, v) for k, v in rule_rows] or ["| (nothing removed) | 0 |"]
    lines += ["", "## Conversations", "", "| Conversation | Turns | Tool calls | Outcome | Flow | Ground truth | Chars before | Chars after |", "|---|---|---|---|---|---|---|---|"]
    for conv in conversations:
        summary = conv.get("summary", {})
        lines.append("| %s | %s | %s | %s | %s | %s | %s | %s |" % (
            conv["conversation_id"], summary.get("turns"), summary.get("tool_calls"), summary.get("outcome"),
            conv["metadata"].get("flow") or "", "yes" if conv["metadata"].get("ground_truth") else "no",
            summary.get("chars_before", ""), summary.get("chars_after", "")))
    return "\n".join(lines) + "\n"


# ----------------------------------------------------------------------------------------------------------
# commands
# ----------------------------------------------------------------------------------------------------------
def detect_format(path: Path, forced: str = "auto") -> str:
    if forced != "auto":
        return forced
    if path.is_dir():
        if list(path.rglob("events.jsonl")):
            return "events_jsonl"
        if any("response" in p.name.lower() for p in path.rglob("*.json")):
            return "ces_json"
        return "unknown"
    suffix = path.suffix.lower()
    if suffix == ".jsonl":
        return "events_jsonl"
    if suffix in (".xlsx", ".xlsm"):
        return "xlsx"
    if suffix in (".csv", ".tsv"):
        return "csv"
    if suffix == ".json":
        return "ces_json"
    return "unknown"


def cmd_inspect(args: argparse.Namespace) -> int:
    path = Path(args.file)
    if not path.exists():
        print("error: %s does not exist" % path, file=sys.stderr)
        return 2
    header, records, _, sheet_name = read_table(path, args.sheet)
    proposal = propose_column_map(header, records)
    proposal["file"] = str(path)
    proposal["sheet"] = sheet_name
    proposal["rows_total"] = len(records)
    proposal["suggested_column_map"]["sheet"] = sheet_name
    out = Path(args.out) if args.out else path.with_name("column_map.proposed.json")
    write_json(out, proposal)
    print("proposal written to %s" % out)
    print("fields: " + ", ".join("%s->%s (%.2f)" % (f, p["column"], p["confidence"]) for f, p in proposal["proposal"].items()))
    if proposal["required_missing_or_low_confidence"]:
        print("needs confirmation for: " + ", ".join(proposal["required_missing_or_low_confidence"]))
    print("confirm the mapping with the user, save it as column_map.json, then run normalize --column-map")
    return 0


def cmd_normalize(args: argparse.Namespace) -> int:
    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    column_map = read_json(args.column_map) if args.column_map else None
    conversations: List[dict] = []
    sources: List[dict] = []
    for raw in args.inputs:
        path = Path(raw)
        if not path.exists():
            print("error: %s does not exist" % path, file=sys.stderr)
            return 2
        fmt = detect_format(path, args.format)
        if fmt == "events_jsonl":
            loaded = []
            for file in discover_events_files(path):
                loaded.extend(load_events_file(file))
        elif fmt == "ces_json":
            loaded = load_ces_json(path)
        elif fmt in ("xlsx", "csv"):
            if column_map is None:
                print("error: %s is a tabular export; run `inspect`, confirm the column map with the user, then pass --column-map" % path, file=sys.stderr)
                return 3
            loaded = load_tabular(path, column_map, args.sheet, fmt)
        else:
            print("error: cannot detect the format of %s (use --format)" % path, file=sys.stderr)
            return 2
        sources.append({"path": str(path), "format": fmt, "conversations": len(loaded)})
        conversations.extend(loaded)
    attached = 0
    if args.ground_truth:
        attached = attach_ground_truth(conversations, Path(args.ground_truth))
    seen = Counter(c["conversation_id"] for c in conversations)
    duplicates = [cid for cid, count in seen.items() if count > 1]
    for conv in conversations:
        if conv["conversation_id"] in duplicates:
            warn(conv, "duplicate_conversation_id", "the same id appears in more than one source")
    diet_settings = diet_settings_from_args(args)
    counts: Dict[str, int] = {}
    for conv in conversations:
        apply_diet(conv, diet_settings, counts)
    diet = {"settings": diet_settings, "counts": counts, "chars_before": counts.get("chars_before", 0), "chars_after": counts.get("chars_after", 0)}
    payload = {"meta": {"skill_version": SKILL_VERSION, "generated_at": now_iso(), "sources": sources,
                        "column_map": column_map, "ground_truth_file": args.ground_truth, "ground_truth_attached": attached,
                        "conversations": len(conversations), "diet": diet},
               "conversations": conversations}
    write_json(out_dir / "conversations.json", payload)
    (out_dir / "normalize_report.md").write_text(build_report(sources, conversations, diet), encoding="utf-8")
    print("normalized %d conversations from %d source(s) -> %s" % (len(conversations), len(sources), out_dir / "conversations.json"))
    redacted = sum(v for k, v in counts.items() if k not in ("chars_before", "chars_after", "truncated", "capped", "snapshots_dropped", "opaque_seen", "old_summarized"))
    print("diet: redacted=%d truncated=%d capped=%d old_summarized=%d snapshots_dropped=%d opaque_seen=%d chars %d -> %d" % (
        redacted, counts.get("truncated", 0), counts.get("capped", 0), counts.get("old_summarized", 0), counts.get("snapshots_dropped", 0),
        counts.get("opaque_seen", 0), counts.get("chars_before", 0), counts.get("chars_after", 0)))
    return 0


def cmd_split(args: argparse.Namespace) -> int:
    data = read_json(args.conversations)
    conversations = data.get("conversations", [])
    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    batches: List[List[dict]] = []
    current: List[dict] = []
    current_chars = 0
    for conv in conversations:
        size = len(json.dumps(conv, ensure_ascii=False))
        if current and (len(current) >= args.max_conversations or current_chars + size > args.max_chars):
            batches.append(current)
            current, current_chars = [], 0
        current.append(conv)
        current_chars += size
    if current:
        batches.append(current)
    oversized = [(c["conversation_id"], len(json.dumps(c, ensure_ascii=False))) for c in conversations
                 if len(json.dumps(c, ensure_ascii=False)) > args.max_chars]
    for cid, size in oversized:
        print("oversized: %s %d chars > %d (one conversation per batch; consider a stricter diet)" % (cid, size, args.max_chars))
    index = []
    for number, batch in enumerate(batches, start=1):
        batch_id = "%03d" % number
        payload = {"meta": dict(data.get("meta", {}), batch_of=len(batches)),
                   "batch": {"id": batch_id, "index": number, "of": len(batches),
                             "conversation_ids": [c["conversation_id"] for c in batch]},
                   "conversations": batch}
        file = out_dir / ("batch_%s.json" % batch_id)
        write_json(file, payload)
        index.append({"id": batch_id, "file": str(file), "conversations": len(batch),
                      "chars": sum(len(json.dumps(c, ensure_ascii=False)) for c in batch)})
    write_json(out_dir / "index.json", {"generated_at": now_iso(), "batches": index, "source": str(args.conversations),
                                        "max_conversations": args.max_conversations, "max_chars": args.max_chars,
                                        "oversized": [{"conversation_id": cid, "chars": size} for cid, size in oversized]})
    print("wrote %d batch file(s) to %s" % (len(batches), out_dir))
    return 0


def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="command", required=True)
    p_inspect = sub.add_parser("inspect", help="propose a column map for an xlsx/csv export")
    p_inspect.add_argument("file")
    p_inspect.add_argument("--sheet")
    p_inspect.add_argument("--out")
    p_inspect.set_defaults(func=cmd_inspect)
    p_norm = sub.add_parser("normalize", help="write conversations.json")
    p_norm.add_argument("inputs", nargs="+")
    p_norm.add_argument("--out", required=True)
    p_norm.add_argument("--column-map")
    p_norm.add_argument("--format", default="auto", choices=["auto", "events_jsonl", "ces_json", "xlsx", "csv"])
    p_norm.add_argument("--ground-truth")
    p_norm.add_argument("--sheet")
    p_norm.set_defaults(func=cmd_normalize)
    p_split = sub.add_parser("split", help="split conversations.json into batches")
    p_split.add_argument("conversations")
    p_split.add_argument("--out", required=True)
    p_split.add_argument("--max-conversations", type=int, default=12)
    p_split.add_argument("--max-chars", type=int, default=40000)
    p_norm.add_argument("--state-snapshots", choices=["none", "turn", "all"], default=DIET_DEFAULTS["state_snapshots"],
                             help="keep the full-state copy on state_update events: none (default; evaluators apply changes), turn (last of each turn), all")
    p_norm.add_argument("--no-redact", action="store_true", help="do not redact bearer tokens, JWTs, secret-named keys and card numbers")
    p_norm.add_argument("--redact-opaque", action="store_true", help="also redact long opaque alphanumeric strings (off: identifiers look the same)")
    p_norm.add_argument("--max-string", type=int, default=DIET_DEFAULTS["max_string"], help="truncate payload strings beyond this many characters (0 = never)")
    p_norm.add_argument("--max-depth", type=int, default=DIET_DEFAULTS["max_depth"], help="payload nesting kept before a container becomes a summary")
    p_norm.add_argument("--list-items", type=int, default=DIET_DEFAULTS["list_items"], help="items kept from a long list (0 = none)")
    p_norm.add_argument("--max-payload", type=int, default=DIET_DEFAULTS["max_payload"], help="hard cap per payload in characters (0 = none)")
    p_norm.add_argument("--no-diet", action="store_true", help="keep everything: all snapshots, no redaction, no truncation (debugging only)")
    p_split.set_defaults(func=cmd_split)
    args = parser.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
